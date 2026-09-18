"""Modo banca do Rubric AI Masters.

Avalia um TCF pronto, sem projeto de capstone, e gera o parecer padronizado da Must
preenchido com nota sugerida e justificativa escolhida do banco de frases das professoras.

O trabalho comentado continua sendo gerado pelo processor.py. Este modulo cuida do parecer.
"""

import json
import logging
import os
from datetime import datetime

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

logger = logging.getLogger("rubric-ai-masters")

MODELO = "gpt-4o"

# Criterios que a IA avalia lendo o texto. Cada frase vem do banco das professoras e ja
# esta amarrada a uma nota. A IA escolhe o indice, nunca escreve a justificativa.
CRITERIOS_PARECER = [
    {
        "chave": "relevancia",
        "titulo": "RELEVÂNCIA E CONTRIBUIÇÃO DO TEMA PARA A ÁREA",
        "itens": ["Objetivos claros", "Justificativa relevante", "Adequação do tema ao conteúdo"],
        "banco": [
            (10.0, "Os objetivos gerais e específicos estão bem formulados e são coerentes com a "
                   "problemática. Há clareza na delimitação temática e nos propósitos da pesquisa."),
            (9.5, "Os objetivos são claros, adequados ao conteúdo e abordam uma temática relevante."),
            (9.0, "A pesquisa realizada tem objetivos claros e apresenta-se como relevante para o "
                  "campo educacional."),
            (8.5, "Objetivos claros, tema relevante, embora amplamente abordado."),
        ],
    },
    {
        "chave": "fundamentacao",
        "titulo": "FUNDAMENTAÇÃO TEÓRICA",
        "itens": ["Consistente, trazendo os principais autores da área",
                  "Atualizada e possui coerência com os objetivos do trabalho"],
        "banco": [
            (10.0, "A fundamentação teórica é atualizada e consistente, demonstrando coerência com "
                   "os objetivos propostos."),
            (9.0, "Boa fundamentação teórica, mas há longos trechos sem apoio teórico para afirmações."),
            (8.5, "Boa fundamentação teórica, mas senti falta de aprofundamento, atualidade e "
                  "diversidade de fontes."),
            (8.5, "Boa fundamentação teórica, mas com várias fontes indicadas no texto ausentes da "
                  "seção de Referências. Recomenda-se uma cuidadosa revisão dos trabalhos utilizados."),
        ],
    },
    {
        "chave": "metodologia",
        "titulo": "METODOLOGIA",
        "itens": ["Descrição metodológica clara", "Tem coerência com os objetivos do trabalho"],
        "banco": [
            (10.0, "A metodologia utilizada demonstra coerência com os objetivos do trabalho, "
                   "apresentando clareza na sua descrição."),
            (9.5, "A metodologia está descrita com clareza e tem coerência com os objetivos do trabalho."),
            (9.0, "Boa descrição metodológica, mas falta dialogar com a literatura de Metodologia "
                  "Científica."),
            (8.5, "Metodologia adequada, mas faltam ajustes conceituais (equívocos conceituais entre "
                  "pesquisa bibliográfica e revisão da literatura)."),
            (8.5, "Descrição metodológica satisfatória."),
            (8.0, "A metodologia necessita de ajustes e mais aprofundamento."),
        ],
    },
    {
        "chave": "pesquisa",
        "titulo": "PESQUISA REALIZADA",
        "itens": ["O problema da pesquisa foi demonstrado no trabalho desenvolvido",
                  "A pesquisa transcorreu de acordo com a metodologia proposta."],
        "banco": [
            (10.0, "A pesquisa transcorreu em conformidade com a metodologia proposta, com o problema "
                   "de pesquisa claramente demonstrado."),
            (9.5, "A pesquisa está bem desenvolvida e fundamentada."),
            (8.5, "O problema de pesquisa foi demonstrado satisfatoriamente."),
            (8.5, "A pesquisa foi desenvolvida conforme o problema de pesquisa, mas faltou "
                  "aprofundamento."),
            (8.0, "O problema de pesquisa, sob a forma de uma pergunta, não foi apresentado, o que "
                  "dificultou a percepção sobre a sua demonstração no decorrer da escrita."),
        ],
    },
    {
        "chave": "textual",
        "titulo": "AVALIAÇÃO TEXTUAL",
        "itens": ["O texto seguiu a estrutura proposta na disciplina Capstone.",
                  "Formatação e organização de acordo com as normas GUIA DE FORMATAÇÃO DO TCF da "
                  "Must University",
                  "Texto de acordo com as normas gramaticais",
                  "As Referências Bibliográficas são citadas no texto e listadas no final do "
                  "trabalho, de acordo com as normas do GUIA DE FORMATAÇÃO DO TCF da Must University"],
        # Decisao da Must: avaliacao textual NUNCA recebe nota 10.
        "banco": [
            (9.7, "O texto está de acordo com a estrutura, formatação e normas gramaticais do Guia de "
                  "Formatação do TCF, necessitando de alguns ajustes de formatação."),
            (9.0, "A redação e a formatação estão adequadas, mas indica-se uma revisão."),
            (8.0, "Indica-se uma revisão cuidadosa de formatação e de ortografia."),
        ],
    },
    {
        "chave": "resultados",
        "titulo": "RESULTADOS E CONCLUSÕES",
        "itens": ["Os resultados estão alinhados com as justificativas e objetivos propostos",
                  "Os resultados são coerentes e relevantes, trazendo contribuições ao tema e "
                  "demonstram análise e resultados coerentes"],
        "banco": [
            (10.0, "Os resultados e conclusões apresentados são coerentes e relevantes, estando "
                   "alinhados aos objetivos e justificativas da pesquisa."),
            (10.0, "As contribuições estão bem sistematizadas e são apresentadas com criticidade."),
            (9.5, "Os resultados estão alinhados e são densos e amplos."),
            (8.5, "Os resultados são relevantes para o campo da pesquisa, mas falta maior alinhamento "
                  "com os objetivos."),
        ],
    },
]

# Criterios de processo. A banca nao acompanhou o aluno, entao ficam em branco.
# O orientador ve estes campos para preencher a mao.
CRITERIOS_PROCESSO = [
    {
        "chave": "autonomia",
        "titulo": "AUTONOMIA DO ALUNO",
        "itens": ["O aluno se envolveu com todo o processo, assumindo responsabilidade pelo projeto.",
                  "O aluno demonstrou independência em todos os aspectos do desenvolvimento, "
                  "implementação e avaliação do projeto."],
    },
    {
        "chave": "criatividade",
        "titulo": "CRIATIVIDADE",
        "itens": ["O aluno demonstrou criatividade, ultrapassando as abordagens convencionais no que "
                  "se refere ao conteúdo do projeto e à metodologia."],
    },
    {
        "chave": "entregas",
        "titulo": "ENTREGA DOS RELATÓRIOS (CAPÍTULOS)",
        "itens": ["Todos as entregas foram concluídas no prazo."],
    },
]


# ------------------------------------------------------------- chamada da IA

def _instrucao_criterios():
    linhas = []
    for c in CRITERIOS_PARECER:
        linhas.append("CRITERIO %s: %s" % (c["chave"], c["titulo"]))
        linhas.append("  O que se avalia: " + " | ".join(c["itens"]))
        linhas.append("  Justificativas disponiveis (escolha UMA pelo indice):")
        for i, (nota, frase) in enumerate(c["banco"]):
            linhas.append("    [%d] nota %s -> %s" % (i, ("%g" % nota).replace(".", ","), frase))
        linhas.append("")
    return "\n".join(linhas)


PROMPT_PARECER = """Voce e avaliador de banca examinadora de mestrado da Must University.
Norma academica: APA. A data de hoje e %s.

Voce vai avaliar um Trabalho de Conclusao Final (TCF) JA CONCLUIDO e preencher o parecer da banca.

Para CADA criterio abaixo, escolha UMA justificativa da lista, pelo indice. A nota vem junto com a
justificativa escolhida, voce NAO define nota livremente e NAO escreve justificativa propria.
Escolha pela evidencia encontrada no texto, nao pela vontade de agradar.

%s

Escreva tambem:
- "titulo": o titulo do TCF, exatamente como aparece no trabalho.
- "parecer": tres paragrafos de parecer da banca, em portugues brasileiro, tom respeitoso e academico.
  O primeiro situa o tema e o que o trabalho investigou. O segundo comenta fundamentacao, metodologia
  e resultados, com base no que voce encontrou. O terceiro fecha reconhecendo o merito do trabalho e
  indicando que ha ajustes a fazer. NAO invente dados que nao estejam no texto.
- "devolutiva": lista de 2 a 5 itens objetivos que o aluno precisa corrigir, cada um comecando por um
  verbo no infinitivo. Aponte apenas o que voce verificou no texto.

Retorne APENAS um JSON valido, sem markdown:
{"criterios": {"relevancia": 0, "fundamentacao": 0, "metodologia": 0, "pesquisa": 0, "textual": 0,
"resultados": 0}, "titulo": "...", "parecer": ["...", "...", "..."], "devolutiva": ["...", "..."]}
Os valores de "criterios" sao os INDICES das justificativas escolhidas."""


def _extrair_json(texto):
    t = (texto or "").strip()
    if "```json" in t:
        t = t.split("```json")[1].split("```")[0]
    elif "```" in t:
        t = t.split("```")[1].split("```")[0]
    return json.loads(t.strip())


def avaliar_para_parecer(cliente, texto_versao, nome_aluno):
    """Uma chamada que devolve nota, justificativa, titulo, parecer e devolutiva."""
    sistema = PROMPT_PARECER % (datetime.now().strftime("%d/%m/%Y"), _instrucao_criterios())
    try:
        resposta = cliente.chat.completions.create(
            model=MODELO,
            messages=[
                {"role": "system", "content": sistema},
                {"role": "user", "content": "TCF de %s:\n\n%s" % (nome_aluno, texto_versao[:90000])},
            ],
            temperature=0.2,
        )
        dados = _extrair_json(resposta.choices[0].message.content)
    except Exception as e:
        logger.error("[parecer] falha na chamada ou no parse: %s" % e)
        raise ValueError("Nao foi possivel gerar o parecer da banca: %s" % e)

    escolhas = dados.get("criterios") or {}
    resultado = {}
    for c in CRITERIOS_PARECER:
        idx = escolhas.get(c["chave"])
        if not isinstance(idx, int) or idx < 0 or idx >= len(c["banco"]):
            logger.warning("[parecer] indice invalido em %s: %r, usando o mais conservador"
                           % (c["chave"], idx))
            idx = len(c["banco"]) - 1
        nota, frase = c["banco"][idx]
        resultado[c["chave"]] = {"nota": nota, "justificativa": frase}
        logger.info("[parecer] %s -> nota %s" % (c["chave"], nota))

    notas = [v["nota"] for v in resultado.values()]
    media = round(sum(notas) / float(len(notas)), 2)
    logger.info("[parecer] media sugerida: %s" % media)

    return {
        "criterios": resultado,
        "media": media,
        "titulo": (dados.get("titulo") or "").strip(),
        "parecer": [p.strip() for p in (dados.get("parecer") or []) if p and p.strip()],
        "devolutiva": [d.strip() for d in (dados.get("devolutiva") or []) if d and d.strip()],
    }


# ------------------------------------------------------- montagem do documento

def _fmt(nota):
    return ("%g" % nota).replace(".", ",")


def _texto(paragrafo, texto, negrito=False, italico=False, tamanho=10, cor=None, realce=False):
    run = paragrafo.add_run(texto)
    run.bold = negrito
    run.italic = italico
    run.font.size = Pt(tamanho)
    run.font.name = "Arial"
    if cor:
        run.font.color.rgb = RGBColor(*cor)
    if realce:
        rpr = run._element.get_or_add_rPr()
        h = OxmlElement("w:highlight")
        h.set(qn("w:val"), "yellow")
        rpr.append(h)
    return run


def _fixar_larguras(tabela, larguras_dxa):
    """python-docx sozinho nao respeita largura de coluna: precisa do layout fixo e do tblGrid."""
    tbl = tabela._tbl
    tblPr = tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    grid = tbl.find(qn("w:tblGrid"))
    if grid is not None:
        tbl.remove(grid)
    grid = OxmlElement("w:tblGrid")
    for largura in larguras_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(largura))
        grid.append(col)
    tbl.insert(1, grid)
    for linha in tabela.rows:
        for celula, largura in zip(linha.cells, larguras_dxa):
            tcPr = celula._tc.get_or_add_tcPr()
            for antigo in tcPr.findall(qn("w:tcW")):
                tcPr.remove(antigo)
            tcW = OxmlElement("w:tcW")
            tcW.set(qn("w:w"), str(largura))
            tcW.set(qn("w:type"), "dxa")
            tcPr.append(tcW)


def _celula(celula, linhas, **kw):
    celula.text = ""
    p = celula.paragraphs[0]
    for i, linha in enumerate(linhas):
        if i:
            p = celula.add_paragraph()
        _texto(p, linha, **kw)


def gerar_parecer_docx(caminho_saida, dados, avaliacao, papel="banca"):
    """Monta o formulario da Must preenchido. papel: 'banca' ou 'orientador'."""
    doc = Document()
    estilo = doc.styles["Normal"]
    estilo.font.name = "Arial"
    estilo.font.size = Pt(10)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _texto(p, "FORMULÁRIO PARA AVALIAÇÃO DO TCF", negrito=True, tamanho=13)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _texto(p, "PARECER DO ORIENTADOR E DOS PROFESSORES DOUTORES DA BANCA",
           negrito=True, tamanho=11)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _texto(p, "Preenchimento sugerido pela Rubric AI a partir da leitura do trabalho. A decisão e o "
              "lançamento das notas na plataforma da Must são do professor.",
           italico=True, tamanho=8, cor=(0x80, 0x80, 0x80))

    # cabecalho
    cab = doc.add_table(rows=3, cols=2)
    cab.style = "Table Grid"
    cab.autofit = False
    cab.alignment = WD_TABLE_ALIGNMENT.CENTER
    _celula(cab.cell(0, 0), ["Aluno(a): " + dados.get("aluno", "")])
    _celula(cab.cell(0, 1), ["Orientador (a): " + dados.get("orientador", "")])
    _celula(cab.cell(1, 0), ["Data: " + dados.get("data", datetime.now().strftime("%d/%m/%Y"))])
    _celula(cab.cell(1, 1), ["Prof. Avaliador (a): " + dados.get("avaliador", "")])
    _celula(cab.cell(2, 0), ["Programa: " + dados.get("programa", "")])
    _celula(cab.cell(2, 1), [""])

    doc.add_paragraph()

    # criterios
    todos = []
    for c in CRITERIOS_PARECER:
        todos.append((c, avaliacao["criterios"].get(c["chave"])))
    ordem_processo = {"autonomia": 5, "criatividade": 6, "entregas": 7}
    for c in CRITERIOS_PROCESSO:
        todos.insert(ordem_processo[c["chave"]], (c, None))

    tab = doc.add_table(rows=1, cols=3)
    tab.style = "Table Grid"
    tab.autofit = False
    LARGURAS = [5400, 900, 3060]
    cab_tab = tab.rows[0]
    for celula, rotulo in zip(cab_tab.cells,
                              ["CRITÉRIOS PARA AVALIAÇÃO DO TCF", "NOTA", "JUSTIFICATIVA"]):
        _celula(celula, [rotulo], negrito=True)
        celula.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    for criterio, resultado in todos:
        linha = tab.add_row()
        _celula(linha.cells[0], [criterio["titulo"]], negrito=True)
        for item in criterio["itens"]:
            _texto(linha.cells[0].add_paragraph(), item)

        if resultado is None:
            _celula(linha.cells[1], [""])
            _celula(linha.cells[2], ["a preencher pelo professor"],
                    italico=True, cor=(0x80, 0x80, 0x80))
        else:
            linha.cells[1].text = ""
            pn = linha.cells[1].paragraphs[0]
            pn.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _texto(pn, _fmt(resultado["nota"]), negrito=True, realce=True)
            _celula(linha.cells[2], [resultado["justificativa"]])

    _fixar_larguras(tab, LARGURAS)
    _fixar_larguras(cab, [4680, 4680])

    def secao(rotulo, paragrafos):
        p = doc.add_paragraph()
        _texto(p, rotulo, negrito=True)
        for texto in paragrafos:
            _texto(doc.add_paragraph(), texto)

    secao("TÍTULO DO TRABALHO DE CONCLUSÃO FINAL (TCF)", [avaliacao.get("titulo") or ""])
    secao("PARECER DO PROFESSOR ORIENTADOR / MEMBRO DA BANCA", avaliacao.get("parecer") or [])
    secao("NOTA NO AVA", ["Média sugerida dos critérios avaliados: " + _fmt(avaliacao["media"])])
    secao("SUGESTÃO DE DEVOLUTIVA", avaliacao.get("devolutiva") or [])

    doc.add_paragraph()
    tem_devolutiva = bool(avaliacao.get("devolutiva"))
    for texto, marcar in [
        ("APROVADO – Sem restrições - Nota: ______", not tem_devolutiva),
        ("APROVADO – Mediante devolutiva com correções - Nota: %s"
         % (_fmt(avaliacao["media"]) if tem_devolutiva else "______"), tem_devolutiva),
        ("REPROVADO - Nota: ______", False),
    ]:
        _texto(doc.add_paragraph(), texto, realce=marcar)

    doc.add_paragraph()
    _texto(doc.add_paragraph(), "Prof. Dr. Avaliador (a): ______________________________________")
    _texto(doc.add_paragraph(), "Prof. Dr. Orientador (a): ______________________________________")

    doc.save(caminho_saida)
    logger.info("[parecer] documento gerado em %s" % caminho_saida)
    return caminho_saida


# ------------------------------------------------------------ fluxo da banca

async def processar_banca(caminho_tcf, nome_aluno, nome_orientador, programa,
                          nome_avaliador, papel="banca"):
    """Avalia um TCF pronto e devolve os dois documentos da banca.

    Retorna (caminho_comentado, caminho_parecer, avaliacao).
    papel: 'banca' deixa autonomia, criatividade e entregas em branco.
           'orientador' tambem deixa em branco, porque a IA nao observa processo,
           mas o nome vai para o campo Orientador do cabecalho.
    """
    import openai
    import tempfile

    from extrator import extrair_texto_docx
    from processor import processar_documento

    texto = extrair_texto_docx(caminho_tcf)
    if not texto or texto.startswith("Erro ao extrair"):
        raise ValueError("Nao foi possivel ler o texto do TCF enviado.")

    # 1. comentarios no trabalho, usando o motor que ja existe, com todos os capitulos
    caminho_comentado = await processar_documento(
        caminho_versao=caminho_tcf,
        caminho_projeto=None,
        nome_aluno=nome_aluno,
        numero_versao="banca",
        capitulos=["introducao", "metodologia", "referencial", "resultados", "conclusao"],
        nome_professor=nome_avaliador or "Membro da Banca",
    )

    # 2. avaliacao dos criterios e textos do parecer
    cliente = openai.OpenAI(api_key=os.environ.get("OPENAI_API_KEY"))
    avaliacao = avaliar_para_parecer(cliente, texto, nome_aluno)

    # 3. montagem do parecer
    with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as tmp:
        caminho_parecer = tmp.name

    dados = {
        "aluno": nome_aluno,
        "orientador": nome_orientador,
        "programa": programa,
        "data": datetime.now().strftime("%d/%m/%Y"),
        "avaliador": nome_avaliador if papel == "banca" else "",
    }
    if papel == "orientador" and nome_avaliador:
        dados["orientador"] = nome_avaliador

    gerar_parecer_docx(caminho_parecer, dados, avaliacao, papel=papel)
    return caminho_comentado, caminho_parecer, avaliacao


def resumo_para_tela(avaliacao):
    """Devolve as notas num formato simples, para a tela de revisao mostrar."""
    linhas = []
    for c in CRITERIOS_PARECER:
        r = avaliacao["criterios"].get(c["chave"]) or {}
        linhas.append({
            "chave": c["chave"],
            "criterio": c["titulo"],
            "nota": r.get("nota"),
            "justificativa": r.get("justificativa"),
        })
    for c in CRITERIOS_PROCESSO:
        linhas.append({
            "chave": c["chave"],
            "criterio": c["titulo"],
            "nota": None,
            "justificativa": "a preencher pelo professor",
        })
    return {
        "criterios": linhas,
        "media": avaliacao.get("media"),
        "titulo": avaliacao.get("titulo"),
        "parecer": avaliacao.get("parecer"),
        "devolutiva": avaliacao.get("devolutiva"),
    }
