from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn

# Quanto texto de UMA tabela vai junto com o paragrafo que a antecede.
#
# Era 4000 e isso causava um erro silencioso e grave: o Quadro das obras do corpus,
# em trabalhos com muitas obras, chegava cortado na analise. Num caso real o quadro
# tinha 4133 caracteres, as ultimas obras sumiam e a IA ainda comentava que o registro
# do aluno estava "interrompido", porque de fato estava interrompido no que ela recebeu.
#
# O limite por chamada ao modelo e de 60000 caracteres por capitulo, entao 12000 para
# uma tabela e folgado. Quando ainda assim estourar, o corte passa a ser DECLARADO no
# texto, para a IA saber que a falta e do sistema e nao do aluno.
LIMITE_TEXTO_TABELA = 12000
AVISO_TABELA_CORTADA = (" [ATENCAO: a tabela e maior do que isto e foi cortada pelo sistema neste "
                        "ponto. O texto do aluno NAO esta incompleto aqui. NAO comente essa "
                        "interrupcao e NAO avalie o que viria depois dela.]")


def extrair_texto_docx(caminho: str) -> str:
    """Extrai o texto completo de um arquivo .docx, numerado por paragrafo.

    ATENCAO: a numeracao [N] desta funcao e usada para ancorar os comentarios
    no documento. Ela precisa continuar percorrendo apenas doc.paragraphs,
    que e o mesmo conjunto usado por body.findall("w:p") no processor.
    NAO acrescentar leitura de tabelas aqui.
    """
    try:
        doc = Document(caminho)
        paragrafos = []
        for i, p in enumerate(doc.paragraphs):
            if p.text.strip():
                paragrafos.append(f"[{i}] {p.text.strip()}")
        return "\n".join(paragrafos)
    except Exception as e:
        return f"Erro ao extrair texto: {str(e)}"


def extrair_texto_docx_com_tabelas(caminho: str) -> str:
    """Mesma numeracao [N] de extrair_texto_docx, mas com o texto das tabelas.

    O conteudo de cada tabela e anexado a linha do paragrafo que vem logo antes dela,
    marcado como (TABELA: ...). A numeracao continua contando apenas os paragrafos do
    corpo, na mesma ordem de doc.paragraphs, entao a ancoragem dos comentarios nao muda.
    Serve para a IA enxergar, por exemplo, a tabela do funil de selecao das obras.
    """
    try:
        doc = Document(caminho)
        linhas = {}
        ordem = []
        indice = -1
        ultimo_com_texto = None
        for bloco in _iterar_blocos(doc.element.body, doc):
            if isinstance(bloco, Paragraph):
                indice += 1
                texto = bloco.text.strip()
                if texto:
                    if _e_recuado(bloco):
                        texto = "(RECUO) " + texto
                    linhas[indice] = texto
                    ordem.append(indice)
                    ultimo_com_texto = indice
            else:
                conteudo = _texto_tabela(bloco)
                if conteudo and ultimo_com_texto is not None:
                    if len(conteudo) > LIMITE_TEXTO_TABELA:
                        conteudo = conteudo[:LIMITE_TEXTO_TABELA] + AVISO_TABELA_CORTADA
                    linhas[ultimo_com_texto] += " (TABELA: %s)" % conteudo
        return "\n".join("[%d] %s" % (i, linhas[i]) for i in ordem)
    except Exception as e:
        return f"Erro ao extrair texto: {str(e)}"


def _e_recuado(paragrafo):
    """Paragrafo recuado da margem, como manda a APA para citacao longa (1,27 cm)."""
    try:
        recuo = paragrafo.paragraph_format.left_indent
        return bool(recuo) and recuo.cm >= 1.0
    except Exception:
        return False


def _texto_tabela(tabela):
    partes = []
    for linha_tabela in tabela.rows:
        celulas = []
        for celula in linha_tabela.cells:
            texto_celula = " ".join(p.text.strip() for p in celula.paragraphs if p.text.strip())
            if texto_celula and (not celulas or celulas[-1] != texto_celula):
                celulas.append(texto_celula)
        if celulas:
            partes.append(" | ".join(celulas))
    return " / ".join(partes)


def extrair_texto_docx_completo(caminho: str) -> str:
    """Extrai o texto de um .docx na ordem real do documento, INCLUINDO o
    conteudo das celulas de tabela.

    Sem numeracao de paragrafo. Usar somente em documentos de referencia,
    como o projeto de capstone aprovado, que nao recebe comentarios.
    """
    try:
        doc = Document(caminho)
        linhas = []
        for bloco in _iterar_blocos(doc.element.body, doc):
            if isinstance(bloco, Paragraph):
                texto = bloco.text.strip()
                if texto:
                    linhas.append(texto)
            else:
                for linha_tabela in bloco.rows:
                    celulas = []
                    for celula in linha_tabela.cells:
                        texto_celula = " ".join(
                            p.text.strip()
                            for p in celula.paragraphs
                            if p.text.strip()
                        )
                        if not texto_celula:
                            continue
                        if celulas and celulas[-1] == texto_celula:
                            continue
                        celulas.append(texto_celula)
                    if celulas:
                        linhas.append(" | ".join(celulas))
        return "\n".join(linhas)
    except Exception as e:
        return f"Erro ao extrair texto: {str(e)}"


def _iterar_blocos(elemento_pai, doc):
    """Percorre paragrafos e tabelas na ordem em que aparecem no documento."""
    for filho in elemento_pai.iterchildren():
        if filho.tag == qn("w:p"):
            yield Paragraph(filho, doc)
        elif filho.tag == qn("w:tbl"):
            yield Table(filho, doc)
